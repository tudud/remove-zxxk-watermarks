from io import BytesIO
from pathlib import Path
import zipfile
from lxml import etree as ET
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from handout_cleaner.models import Options, Record
from handout_cleaner.ocr import OCR
from handout_cleaner.slides import clean_pptx


def deck():
    p=Presentation();p.slide_width=Inches(14);p.slide_height=Inches(8)
    return p


def add(slide,text,x=.5,y=1,w=12,h=.5,red=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    shape.text=text
    if red:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb=RGBColor(192,0,0);run.font.size=Pt(28);run.font.bold=True
    return shape


def clean(tmp_path,p,**kwargs):
    source=tmp_path/'in.pptx';target=tmp_path/'out.pptx';p.save(source)
    record=Record(str(source))
    clean_pptx(source,target,Options(ocr=False,remove_last_slide=False,**kwargs),record,OCR())
    return Presentation(target),record,target


def all_text(p):
    return '\n'.join(shape.text for slide in p.slides for shape in slide.shapes if shape.has_text_frame)


def test_split_runs_and_color_are_preserved(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    add(s,'1．选择正确选项。')
    shape=add(s,'',y=2)
    paragraph=shape.text_frame.paragraphs[0]
    paragraph.add_run().text='【答';paragraph.add_run().text='案】'
    run=paragraph.add_run();run.text='Ｂ';run.font.bold=True;run.font.color.rgb=RGBColor(192,0,0)
    add(s,'【解析】计算过程保留，故选A。',y=3)
    result,r,_=clean(tmp_path,p)
    assert '【答案】A' in all_text(result) and len(r.changes)==1
    assert r.changes[0]['slide']==1 and r.changes[0]['question']=='1'
    corrected=result.slides[0].shapes[1].text_frame.paragraphs[0].runs[-1]
    assert corrected.font.bold and str(corrected.font.color.rgb)=='C00000'
    assert result.slides[0].shapes[1].left==shape.left


def test_standalone_red_letter_overlay(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    add(s,'【典例3】下列说法正确的是（    ）',w=10)
    add(s,'B',x=9,y=1,w=.5,red=True)
    add(s,'A．甲  B．乙  C．丙  D．丁',y=2)
    add(s,'【解析】故选A。',y=3)
    result,r,_=clean(tmp_path,p)
    assert result.slides[0].shapes[1].text=='A' and len(r.changes)==1
    assert 'A．甲  B．乙' in all_text(result)


def test_inline_answer_not_diagram_label(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    add(s,'【例题2】如图（A），下列错误的是（Ｂ）')
    add(s,'【解析】故选C。',y=2)
    result,r,_=clean(tmp_path,p)
    assert '如图（A），下列错误的是（C）' in all_text(result)
    assert len(r.changes)==1


def test_two_vertical_questions_do_not_mix_conclusions(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    for num,y,key,final in ((1,1,'B','A'),(2,4,'D','C')):
        add(s,f'{num}．选择正确选项。',y=y)
        add(s,'【答案】'+key,y=y+.8)
        add(s,'【解析】故选'+final+'。',y=y+1.5)
    result,r,_=clean(tmp_path,p)
    assert '【答案】A' in all_text(result) and '【答案】C' in all_text(result)
    assert len(r.changes)==2


def test_two_columns_are_scoped_separately(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    for num,x,key,final in ((1,.5,'B','A'),(2,7,'D','C')):
        add(s,f'{num}．选择正确选项。',x=x,w=5.5)
        add(s,'【答案】'+key,x=x,y=2,w=5.5)
        add(s,'【解析】故选'+final+'。',x=x,y=3,w=5.5)
    result,r,_=clean(tmp_path,p)
    assert [c['after'] for c in r.changes]==['A','C']


def test_multiple_questions_in_one_textbox(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    add(s,'1．问题一\n【答案】B\n【解析】故选A。\n2．问题二\n【答案】D\n【解析】故选C。',h=6)
    result,r,_=clean(tmp_path,p)
    assert len(r.changes)==2 and '【答案】A' in all_text(result) and '【答案】C' in all_text(result)


def test_explicit_next_page_explanation(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    add(s,'【典例1】选择正确选项。');add(s,'【答案】B',y=2)
    s=p.slides.add_slide(p.slide_layouts[6]);add(s,'【解析】推导后，故选A。')
    result,r,_=clean(tmp_path,p)
    assert '【答案】A' in all_text(result) and r.changes[0]['slide']==1


def test_repeated_numbers_on_different_pages_do_not_cross(tmp_path):
    p=deck()
    for question,explanation in (('第一道题','没有明确结论。'),('另一道题','故选A。')):
        s=p.slides.add_slide(p.slide_layouts[6]);add(s,'【例题1】'+question)
        add(s,'【答案】B',y=2);add(s,'【解析】'+explanation,y=3)
    result,r,_=clean(tmp_path,p)
    assert result.slides[0].shapes[1].text=='【答案】B'
    assert result.slides[1].shapes[1].text=='【答案】A'
    assert len(r.changes)==1 and r.changes[0]['slide']==2 and r.warnings


def test_conflicting_and_missing_conclusions_are_not_guessed(tmp_path):
    p=deck()
    for explanation in ('故选A。因此选C。','A正确，B错误，但没有显式最终选项。'):
        s=p.slides.add_slide(p.slide_layouts[6]);add(s,'1．问题')
        add(s,'【答案】B',y=2);add(s,'【解析】'+explanation,y=3)
    result,r,_=clean(tmp_path,p)
    assert not r.changes and len(r.warnings)==2 and all_text(result).count('【答案】B')==2


def test_toggle_off_keeps_original_key(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    add(s,'1．问题\n【答案】B\n【解析】故选A。',h=4)
    result,r,_=clean(tmp_path,p,correct_answers=False)
    assert '【答案】B' in all_text(result) and not r.changes


def test_grouped_text_shapes(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6]);g=s.shapes.add_group_shape()
    for text,y in (('【例题4】问题',1),('【答案】B',2),('【解析】故选A。',3)):
        shape=g.shapes.add_textbox(Inches(1),Inches(y),Inches(10),Inches(.5));shape.text=text
    result,r,_=clean(tmp_path,p)
    assert r.changes[0]['after']=='A'
    assert result.slides[0].shapes[0].shapes[1].text=='【答案】A'


def test_table_key_label_and_value(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6]);add(s,'1．问题')
    table=s.shapes.add_table(1,2,Inches(1),Inches(2),Inches(8),Inches(.5)).table
    table.cell(0,0).text='【答案】';table.cell(0,1).text='B'
    add(s,'【解析】故选A。',y=3)
    result,r,_=clean(tmp_path,p)
    assert result.slides[0].shapes[1].table.cell(0,1).text=='A' and len(r.changes)==1


def test_numbered_answer_summary_and_ambiguous_number(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    add(s,'1．问题\n【答案】B\n【解析】故选A。',h=4)
    s=p.slides.add_slide(p.slide_layouts[6]);add(s,'参考答案\n1. B',h=2)
    result,r,_=clean(tmp_path,p)
    assert '1. A' in all_text(result) and len(r.changes)==2


def test_image_only_answer_is_reported_not_rewritten(tmp_path):
    class StubOCR:
        def is_logo(self,data):return False
        def clean_image(self,data):return data,0
        def recognize_image(self,data):return [([], '【例题1】答案：B，故选A。', .99)]
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    data=BytesIO();Image.new('RGB',(500,500),'white').save(data,format='PNG');data.seek(0)
    s.shapes.add_picture(data,Inches(1),Inches(1),Inches(8),Inches(6))
    source=tmp_path/'in.pptx';out=tmp_path/'out.pptx';p.save(source)
    r=Record(str(source));clean_pptx(source,out,Options(remove_last_slide=False),r,StubOCR())
    assert not r.changes and any('图片' in w for w in r.warnings)
    assert Presentation(out).slides[0].shapes[0].image.blob==Presentation(source).slides[0].shapes[0].image.blob


def test_multiselect_normalization(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    add(s,'1．选择正确选项。\n【答案】Ｃ、Ｄ\n【解析】故选BA。',h=4)
    result,r,_=clean(tmp_path,p)
    assert '【答案】AB' in all_text(result) and r.changes[0]['after']=='AB'


def test_summary_does_not_merge_reused_numbers(tmp_path):
    p=deck()
    for stem,final in (('题目一','A'),('题目二','C')):
        s=p.slides.add_slide(p.slide_layouts[6]);add(s,'【例题1】'+stem+'\n【答案】B\n【解析】故选'+final+'。',h=4)
    s=p.slides.add_slide(p.slide_layouts[6]);add(s,'参考答案\n1. B',h=2)
    result,r,_=clean(tmp_path,p)
    assert '1. B' in result.slides[2].shapes[0].text and any('无法唯一关联' in w for w in r.warnings)


def test_absolute_package_targets_support_correction_and_slide_removal(tmp_path):
    p=deck();s=p.slides.add_slide(p.slide_layouts[6])
    add(s,'1．问题\n【答案】B\n【解析】故选A。',h=4)
    s=p.slides.add_slide(p.slide_layouts[6]);add(s,'结束页')
    buffer=BytesIO();p.save(buffer)
    source=tmp_path/'absolute.pptx';target=tmp_path/'out.pptx'
    with zipfile.ZipFile(buffer) as original, zipfile.ZipFile(source,'w') as rewritten:
        for name in original.namelist():
            data=original.read(name)
            if name=='ppt/_rels/presentation.xml.rels':
                root=ET.fromstring(data)
                for relation in root:
                    if relation.get('Type','').endswith('/slide'):
                        relation.set('Target','/ppt/'+relation.get('Target'))
                data=ET.tostring(root)
            rewritten.writestr(name,data)
    record=Record(str(source))
    clean_pptx(source,target,Options(ocr=False),record,OCR())
    result=Presentation(target)
    assert len(result.slides)==1 and '【答案】A' in all_text(result)
    assert len(record.changes)==1 and record.changes[0]['slide']==1
