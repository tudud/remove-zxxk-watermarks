from pptx import Presentation
from pptx.util import Inches
from handout_cleaner.slides import clean_pptx
from handout_cleaner.models import Options, Record
from handout_cleaner.ocr import OCR
from PIL import Image, ImageDraw, ImageFont
from io import BytesIO
import zipfile
from lxml import etree as ET


def test_ppt_source_and_mixed_watermark_preserve_question(tmp_path):
    source=tmp_path/'in.pptx';out=tmp_path/'out.pptx'
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6])
    box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(2))
    box.text_frame.text='1．（2024年临平区一模）题干继续保留。 学科网'
    box=slide.shapes.add_textbox(Inches(1),Inches(4),Inches(4),Inches(1))
    box.text_frame.text='www.zxxk.com'
    p.save(source)
    clean_pptx(source,out,Options(ocr=False),Record(str(source)),OCR())
    result=Presentation(out)
    text='\n'.join(s.text for s in result.slides[0].shapes if s.has_text_frame)
    assert '题干继续保留' in text and '临平区' not in text
    assert '学科网' not in text and 'zxxk' not in text


def test_transparent_logo_in_group_and_last_slide(tmp_path):
    logo = Image.new('RGBA', (400,90))
    font = ImageFont.truetype('vendor/fonts/NotoSansCJKsc-Regular.otf', 50)
    ImageDraw.Draw(logo).text((2,0),'学科网·上好课',font=font,fill='black')
    stream=BytesIO();logo.save(stream,format='PNG');stream.seek(0)
    p=Presentation()
    for i in range(3):
        slide=p.slides.add_slide(p.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(2),Inches(4),Inches(1)).text=f'Lesson {i}'
        group=slide.shapes.add_group_shape()
        stream.seek(0)
        group.shapes.add_picture(stream,Inches(6),Inches(.2),Inches(2),Inches(.45))
    p.slides[-1].notes_slide.notes_text_frame.text='Ending notes'
    source=tmp_path/'in.pptx';out=tmp_path/'out.pptx';p.save(source)
    record=Record(str(source))
    clean_pptx(source,out,Options(),record,OCR())
    result=Presentation(out)
    assert len(result.slides)==2 and record.actions
    assert all(not list(shape.shapes) for slide in result.slides for shape in slide.shapes if shape.shape_type==6)
    with zipfile.ZipFile(out) as z:
        assert 'ppt/slides/slide3.xml' not in z.namelist()
        assert 'ppt/notesSlides/notesSlide1.xml' not in z.namelist()
        assert b'/ppt/slides/slide3.xml' not in z.read('[Content_Types].xml')


def test_keep_last_slide_option(tmp_path):
    p=Presentation()
    for _ in range(2):p.slides.add_slide(p.slide_layouts[6])
    source=tmp_path/'in.pptx';out=tmp_path/'out.pptx';p.save(source)
    clean_pptx(source,out,Options(ocr=False,remove_last_slide=False),Record(str(source)),OCR())
    assert len(Presentation(out).slides)==2
