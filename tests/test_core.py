import zipfile
from pathlib import Path
from lxml import etree as ET
from docx import Document
from handout_cleaner.rules import source_edits, apply_edits, clean_name
from handout_cleaner.word import clean_docx, text_of, W, has_answers
from handout_cleaner.models import Options, Record
from handout_cleaner.archives import discover


def content(path):
    return '\n'.join(p.text for p in Document(path).paragraphs)


def test_source_modes_and_normal_years():
    s='1．（2024年临平区一模）某物体于2024年开始运动。'
    assert apply_edits(s, source_edits(s,'all'))=='1．某物体于2024年开始运动。'
    assert apply_edits(s, source_edits(s,'year'))=='1．（临平区一模）某物体于2024年开始运动。'
    assert not source_edits('正文（2024年临平区一模）补充说明')
    assert not source_edits('2．（2024年）某年温度')
    assert not source_edits('（期末复习）第1章 对环境的察觉')
    assert apply_edits('3.（杭州中考）题干',source_edits('3.（杭州中考）题干'))=='3.题干'


def test_names():
    assert clean_name('第2节+物质的导电性+第1课时（分层作业）科学浙教版2024八年级上册.docx')=='第2节 物质的导电性 第1课时'
    assert clean_name('第4节 欧姆定律及其应用（第3课时）（分层作业）科学浙教版2024八年级上册（答案）.docx')=='第4节 欧姆定律及其应用（第3课时）'
    assert '重难点训练' in clean_name('第1节（重难点训练）科学新教材浙教版九年级上册.docx')


def test_word_correction_preserves_math_and_images_and_splits(tmp_path):
    source=tmp_path/'解析版.docx';answer=tmp_path/'答案.docx';practice=tmp_path/'习题.docx'
    d=Document();d.add_paragraph('来源测试（解析版）')
    p=d.add_paragraph();p.add_run('1．（2024年临');p.add_run('平区一模）选择正确选项。')
    d.add_paragraph('A.甲 B.乙 C.丙 D.丁')
    d.add_paragraph('【答案】B')
    d.add_paragraph('【解析】计算 x=2，故选A。')
    d.add_paragraph('2．第二题。')
    d.add_paragraph('【答案】C')
    d.add_paragraph('【解析】此段没有明确最终选项，不能猜测。')
    d._element.body.append(ET.Comment('real-world documents may include XML comments'))
    d.sections[0].header.paragraphs[0].text='学科网 www.zxxk.com'
    d.sections[0].footer.paragraphs[0].text='学科网'
    d.save(source)
    record=Record(str(source))
    clean_docx(source,answer,Options(ocr=False),record)
    t=content(answer)
    assert '【答案】A' in t
    assert '【答案】C' in t
    assert '临平区一模' not in t
    assert '计算 x=2，故选A。' in t
    assert len(record.changes)==1
    clean_docx(answer,practice,Options(ocr=False),Record(str(answer)),practice=True,remove_answers=True)
    t=content(practice)
    assert '【答案】' not in t and '【解析】' not in t
    assert '第二题' in t and '选择正确选项' in t
    assert '解析版' not in t
    assert not has_answers(practice)
    with zipfile.ZipFile(answer) as z:
        assert b'PAGE' in z.read('word/cleaner-footer.xml')
        assert '学科网' not in z.read('word/header1.xml').decode()


def test_ambiguous_conclusions_do_not_correct(tmp_path):
    source=tmp_path/'in.docx';out=tmp_path/'out.docx'
    d=Document();d.add_paragraph('1．题目');d.add_paragraph('【答案】B');d.add_paragraph('【解析】故选A。因此选C。');d.save(source)
    r=Record(str(source));clean_docx(source,out,Options(ocr=False),r)
    assert '【答案】B' in content(out) and not r.changes and r.warnings


def test_zip_slip_rejected(tmp_path):
    bad=tmp_path/'bad.zip'
    with zipfile.ZipFile(bad,'w') as z:z.writestr('../escape.docx','bad')
    import pytest
    with pytest.raises(ValueError):discover([bad],tmp_path/'unpack')


def test_nested_zip_long_names(tmp_path):
    inner=tmp_path/'inner.zip'
    with zipfile.ZipFile(inner,'w') as z:z.writestr('课程资料/'+'长文件名'*15+'.docx',b'test')
    outer=tmp_path/'outer.zip'
    with zipfile.ZipFile(outer,'w') as z:z.write(inner,'资料/inner.zip')
    result=discover([outer],tmp_path/'unpack')
    assert len(result)==1 and result[0][0].suffix=='.docx'


def test_answer_cells_remain_valid_word_cells(tmp_path):
    source=tmp_path/'table.docx';out=tmp_path/'practice.docx'
    d=Document();t=d.add_table(rows=2,cols=2)
    t.cell(0,0).text='1．题目甲';t.cell(0,1).text='【答案】B'
    t.cell(1,0).text='2．题目乙';t.cell(1,1).text='【答案】C'
    d.save(source)
    clean_docx(source,out,Options(ocr=False),Record(str(source)),practice=True,remove_answers=True)
    result=Document(out)
    assert len(result.tables)==1
    assert all(cell.paragraphs for row in result.tables[0].rows for cell in row.cells)
    assert '题目甲' in result.tables[0].cell(0,0).text and not result.tables[0].cell(0,1).text
