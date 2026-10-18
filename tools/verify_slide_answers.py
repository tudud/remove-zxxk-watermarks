"""Check an intentionally mismatched answer in an excerpt of a real deck."""
import argparse
import json
from pathlib import Path
import re
from lxml import etree as ET
from pptx import Presentation
import pymupdf as fitz
from handout_cleaner.convert import convert_office
from validate_portable import execute

ROOT=Path(__file__).resolve().parents[1]
A='http://schemas.openxmlformats.org/drawingml/2006/main'


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--app',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    output=args.output.resolve();output.mkdir(parents=True,exist_ok=True)
    original=next((ROOT/'artifacts/patch-qa/real-inputs').glob('4.1*.pptx'))
    presentation=Presentation(original)
    slide=presentation.slides[22]
    key=next(s for s in slide.shapes if getattr(s,'text','').strip()=='B')
    key._element.find('.//{'+A+'}t').text='C'
    explanation=next(s for s in slide.shapes if '【解析】' in getattr(s,'text',''))
    for index,node in enumerate(explanation._element.iter('{'+A+'}t')):
        node.text='【解析】逐项比较可知，故选B。' if index==0 else ''
    keep={presentation.slides._sldIdLst[22].rId,presentation.slides._sldIdLst[-1].rId}
    for item in list(presentation.slides._sldIdLst):
        if item.rId not in keep:
            presentation.part.drop_rel(item.rId);presentation.slides._sldIdLst.remove(item)
    source=output/'真实模板-答案错配测试.pptx'
    presentation.save(source)
    execute(args.app/'HandoutCleaner.exe',['--cli','--input',source,'--output',output/'processed'],output)
    result=next((output/'processed').glob('*.pptx'))
    cleaned=Presentation(result)
    fixed=next(s for s in cleaned.slides[0].shapes if s.shape_id==key.shape_id)
    assert fixed.text=='B'
    assert (fixed.left,fixed.top,fixed.width,fixed.height)==(key.left,key.top,key.width,key.height)
    before_props=[ET.tostring(n) for n in key._element.iter('{'+A+'}rPr')]
    after_props=[ET.tostring(n) for n in fixed._element.iter('{'+A+'}rPr')]
    assert before_props==after_props
    assert len(cleaned.slides)==1
    report=json.loads((output/'processed/处理报告.json').read_text(encoding='utf-8'))
    assert report['version']=='1.0.3'
    assert [(c['before'],c['after'],c['slide']) for c in report['records'][0]['changes']]==[('C','B',1)]
    for label,path in (('before',source),('after',result)):
        pdf=convert_office(path,output/label,'pdf:impress_pdf_Export')
        with fitz.open(pdf) as doc:doc[0].get_pixmap(matrix=fitz.Matrix(1,1)).save(output/(label+'.png'))
    summary={'real_template_red_C_corrected_to_B':True,'font_color_and_position_preserved':True,'closing_slide_deleted':True,'report_has_slide_and_evidence':True}
    (output/'real-template-validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False))


if __name__=='__main__':main()
