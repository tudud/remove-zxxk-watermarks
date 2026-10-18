import argparse
import ctypes
from ctypes import wintypes
import json
import os
from pathlib import Path
import subprocess
import time
from docx import Document

ROOT=Path(__file__).resolve().parents[1]


def execute(exe, arguments, cwd, timeout=180):
    env=os.environ.copy()
    for key in ('PYTHONPATH','PYTHONHOME','QT_PLUGIN_PATH','QT_QPA_PLATFORM','QML2_IMPORT_PATH'):
        env.pop(key,None)
    env['PATH']=os.environ['SystemRoot']+'\\System32;'+os.environ['SystemRoot']
    completed=subprocess.run([str(exe),*map(str,arguments)],cwd=cwd,env=env,timeout=timeout,creationflags=subprocess.CREATE_NO_WINDOW)
    if completed.returncode:
        raise RuntimeError(f'Frozen EXE exit {completed.returncode}: {arguments}')


def gui_check(exe, output):
    import psutil
    ctypes.windll.user32.SetProcessDPIAware()
    env=os.environ.copy()
    for key in ('PYTHONPATH','PYTHONHOME','QT_PLUGIN_PATH','QT_QPA_PLATFORM'):
        env.pop(key,None)
    p=subprocess.Popen([str(exe)],cwd=output,env=env,creationflags=subprocess.CREATE_NO_WINDOW)
    user=ctypes.windll.user32
    callback_type=ctypes.WINFUNCTYPE(wintypes.BOOL,wintypes.HWND,wintypes.LPARAM)
    user.GetWindowThreadProcessId.argtypes=[wintypes.HWND,ctypes.POINTER(wintypes.DWORD)]
    user.IsWindowVisible.argtypes=[wintypes.HWND]
    user.GetWindowTextW.argtypes=[wintypes.HWND,wintypes.LPWSTR,ctypes.c_int]
    user.PostMessageW.argtypes=[wintypes.HWND,wintypes.UINT,wintypes.WPARAM,wintypes.LPARAM]
    found=[]
    diagnostics=[]
    def callback(hwnd,lparam):
        process=wintypes.DWORD()
        user.GetWindowThreadProcessId(hwnd,ctypes.byref(process))
        if process.value in process_ids and user.IsWindowVisible(hwnd):
            buffer=ctypes.create_unicode_buffer(256)
            user.GetWindowTextW(hwnd,buffer,256)
            if buffer.value and buffer.value not in diagnostics:diagnostics.append(buffer.value)
            if '讲义净化' in buffer.value:found.append(hwnd)
        return True
    cb=callback_type(callback)
    try:
        for _ in range(250):
            process_ids={p.pid}
            try:process_ids.update(c.pid for c in psutil.Process(p.pid).children(recursive=True))
            except psutil.Error:pass
            user.EnumWindows(cb,0)
            if found:break
            if p.poll() is not None:raise RuntimeError('GUI exited before a window was created')
            time.sleep(.2)
        if not found:
            (output/'gui-startup-diagnostics.json').write_text(json.dumps(diagnostics,ensure_ascii=False),encoding='utf-8')
            raise RuntimeError('GUI window not found: '+repr(diagnostics))
        time.sleep(1)
        from PIL import ImageGrab
        rect=wintypes.RECT()
        user.GetWindowRect(found[0],ctypes.byref(rect))
        ImageGrab.grab((rect.left,rect.top,rect.right,rect.bottom)).save(output/'frozen-gui.png')
        user.PostMessageW(found[0],0x0010,0,0)
        p.wait(timeout=20)
    finally:
        if p.poll() is None:p.terminate();p.wait(timeout=10)


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--app',type=Path,default=ROOT/'dist'/'HandoutCleaner')
    parser.add_argument('--output',type=Path,default=ROOT/'artifacts'/'portable-check')
    parser.add_argument('--background-fixture',type=Path)
    args=parser.parse_args()
    exe=(args.app/'HandoutCleaner.exe').resolve()
    output=args.output.resolve();output.mkdir(parents=True,exist_ok=True)
    execute(exe,['--self-check','--output',output/'self-check.json'],output)
    checks=json.loads((output/'self-check.json').read_text())
    assert all(checks.values())
    fixtures=ROOT/'artifacts'/'fixtures'
    execute(exe,['--cli','--input',fixtures/'学科网测试包.zip','--output',output/'word','--no-ocr'],output)
    word=json.loads((output/'word'/'处理报告.json').read_text(encoding='utf-8'))
    assert len(word['records'])==1 and word['records'][0]['status']!='失败'
    assert len(word['records'][0]['outputs'])==3
    assert word['records'][0]['changes'][0]['before']=='B' and word['records'][0]['changes'][0]['after']=='A'
    answer=next((output/'word').glob('*（答案）.docx'))
    text='\n'.join(p.text for p in Document(answer).paragraphs)
    assert '【答案】A' in text and '临平区一模' not in text
    seven=fixtures/'便携7z测试.7z'
    if seven.exists():
        execute(exe,['--cli','--input',seven,'--output',output/'seven','--no-ocr'],output)
        report=json.loads((output/'seven'/'处理报告.json').read_text(encoding='utf-8'))
        assert report['records'][0]['status']!='失败' and len(report['records'][0]['outputs'])==3
    from pptx import Presentation
    from pptx.util import Inches
    slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[6])
    box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(2))
    box.text_frame.text='1．（2024年临平区一模）题干保留。\n【答案】B\n【解析】计算后，故选A。 学科网'
    second=slide.shapes.add_textbox(Inches(1),Inches(4),Inches(7),Inches(.5))
    second.text='【例题2】下列说法正确的是（    ）'
    overlay=slide.shapes.add_textbox(Inches(6),Inches(4),Inches(.5),Inches(.5))
    overlay.text='B'
    from pptx.dml.color import RGBColor
    overlay.text_frame.paragraphs[0].runs[0].font.color.rgb=RGBColor(192,0,0)
    explanation=slide.shapes.add_textbox(Inches(1),Inches(5),Inches(7),Inches(.5))
    explanation.text='【解析】逐项比较，故选C。'
    from PIL import Image, ImageDraw, ImageFont
    logo=Image.new('RGBA',(450,100))
    font=ImageFont.truetype(str(ROOT/'vendor/fonts/NotoSansCJKsc-Regular.otf'),50)
    ImageDraw.Draw(logo).text((2,0),'学科网·上好课',font=font,fill='black')
    logo_path=output/'transparent-logo.png';logo.save(logo_path)
    slide.shapes.add_picture(str(logo_path),Inches(6),Inches(.1),Inches(2),Inches(.45))
    slides.slides.add_slide(slides.slide_layouts[6])
    source_ppt=output/'课件测试.pptx';slides.save(source_ppt)
    execute(exe,['--cli','--input',source_ppt,'--output',output/'ppt'],output)
    ppt=Presentation(next((output/'ppt').glob('*.pptx')))
    ppt_text=''.join(s.text for s in ppt.slides[0].shapes if s.has_text_frame)
    assert '题干保留' in ppt_text and '学科网' not in ppt_text and '临平区' not in ppt_text
    assert len(ppt.slides)==1 and not any(s.shape_type==13 for s in ppt.slides[0].shapes)
    assert '【答案】A' in ppt_text and 'C'==ppt.slides[0].shapes[2].text
    ppt_report=json.loads((output/'ppt'/'处理报告.json').read_text(encoding='utf-8'))
    assert [(c['before'],c['after'],c['slide']) for c in ppt_report['records'][0]['changes']]==[('B','A',1),('B','C',1)]
    execute(exe,['--cli','--input',fixtures/'浙考神墙扫描测试.pdf','--output',output/'scan'],output)
    scan=json.loads((output/'scan'/'处理报告.json').read_text(encoding='utf-8'))
    assert scan['records'][0]['status']!='失败' and len(scan['records'][0]['outputs'])==1
    assert not list((output/'scan').glob('*.docx'))
    execute(exe,['--cli','--input',fixtures/'可搜索PDF测试.pdf','--output',output/'searchable','--no-ocr'],output)
    searchable=json.loads((output/'searchable'/'处理报告.json').read_text(encoding='utf-8'))
    assert searchable['records'][0]['status']!='失败'
    assert len(searchable['records'][0]['outputs'])==1 and not list((output/'searchable').glob('*.docx'))
    if args.background_fixture:
        import pymupdf as fitz
        import numpy as np
        from io import BytesIO
        sample=output/'背景扫描测试.pdf'
        with fitz.open(args.background_fixture) as source:
            first=fitz.open();first.insert_pdf(source,from_page=0,to_page=0);first.save(sample);first.close()
        execute(exe,['--cli','--input',sample,'--output',output/'background'],output)
        cleaned=next((output/'background').glob('*.pdf'))
        with fitz.open(sample) as before,fitz.open(cleaned) as after:
            a=np.asarray(Image.open(BytesIO(before.extract_image(before[0].get_images()[0][0])['image'])).convert('RGB'))
            b=np.asarray(Image.open(BytesIO(after.extract_image(after[0].get_images()[0][0])['image'])).convert('RGB'))
        assert a.shape==b.shape
        dark=a.min(axis=2)<180
        assert np.array_equal(a[dark],b[dark]),'Faint-watermark cleanup changed dark printing'
        assert ((a!=b).any(axis=2)).sum()>10000
    gui_check(exe,output)
    result={'self_check':checks,'word_three_outputs':True,'answer_B_corrected_to_A':True,'nested_7zip_processed':seven.exists(),'ppt_question_preserved':True,'scanned_pdf_only_one_pdf':True,'searchable_pdf_only_one_pdf':True,'gui_started_and_closed':True,'system_python_not_on_PATH':True}
    result.update(ppt_transparent_logo_removed=True,ppt_last_slide_deleted=True,faint_background_preserves_dark_pixels=bool(args.background_fixture))
    result.update(ppt_explicit_answer_corrected=True,ppt_red_overlay_answer_corrected=True)
    (output/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
